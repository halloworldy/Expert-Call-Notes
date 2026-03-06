#!/usr/bin/env python3
"""
Expert Tracker Excel merge script.

Usage:
  python merge.py <master_path> <output_path> [--alphasights <path>] [--guidepoint <path>] [--glg <path>]

Outputs: updated master Excel + JSON change log to stdout.
"""

import openpyxl
import openpyxl.styles
import unicodedata
import re
import json
import sys
from copy import copy
from datetime import datetime


def normalise_name(name):
    """Strip whitespace, lowercase, remove accents."""
    if not name:
        return ""
    name = str(name).strip().lower()
    name = unicodedata.normalize("NFD", name)
    name = "".join(c for c in name if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", name)


def find_header_row(ws):
    """Find the row containing column headers (look for 'Name' in col D-G)."""
    for row in range(1, min(15, ws.max_row + 1)):
        for col in range(4, 8):
            cell = ws.cell(row=row, column=col)
            if cell.value and str(cell.value).strip().lower() == "name":
                return row
    return None


def find_network_columns(ws, header_row):
    """Dynamically find AlphaSights, Guidepoint, GLG, Credits, Name, etc. columns."""
    cols = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col).value
        if val:
            val_clean = str(val).strip().lower()
            if "alphasights" in val_clean:
                cols["AlphaSights"] = col
            elif "guidepoint" in val_clean:
                cols["Guidepoint"] = col
            elif "glg" in val_clean:
                cols["GLG"] = col
            elif "credits" in val_clean:
                cols["Credits"] = col
            elif val_clean == "name":
                cols["Name"] = col
            elif "biography" in val_clean or "bio" in val_clean:
                cols["Biography"] = col
            elif "compan" in val_clean:
                cols["Company"] = col
            elif "position" in val_clean:
                cols["Positions"] = col
            elif "country" in val_clean:
                cols["Country"] = col
            elif "start" in val_clean and "date" in val_clean:
                cols["StartDate"] = col
            elif "end" in val_clean and "date" in val_clean:
                cols["EndDate"] = col
            elif "linkedin" in val_clean:
                cols["LinkedIn"] = col
    return cols


def copy_row_style(ws, source_row, target_row):
    """Copy formatting from source row to target row."""
    for col in range(1, ws.max_column + 1):
        src = ws.cell(row=source_row, column=col)
        tgt = ws.cell(row=target_row, column=col)
        if src.has_style:
            tgt.font = copy(src.font)
            tgt.border = copy(src.border)
            tgt.fill = copy(src.fill)
            tgt.number_format = src.number_format
            tgt.protection = copy(src.protection)
            tgt.alignment = copy(src.alignment)


def _resolve_target_sheet(master_wb, net_sheet_name, company_name, network_name):
    """Map network sheet name to master sheet. Apply routing rules."""
    # Special Guidepoint routing: Telepass/Via Verde from Others -> Telepass tab
    if network_name == "Guidepoint" and "other" in net_sheet_name.lower():
        company_lower = company_name.lower()
        if "telepass" in company_lower or "via verde" in company_lower:
            for ms in master_wb.sheetnames:
                if "telepass" in ms.lower():
                    return ms

    # Standard matching - find master sheet whose title row contains the network sheet name
    net_clean = net_sheet_name.strip().lower()
    for ms in master_wb.sheetnames:
        ws = master_wb[ms]
        title_val = ws.cell(row=1, column=2).value
        if title_val and net_clean in str(title_val).strip().lower():
            return ms

    # Fallback: partial match on sheet name
    for ms in master_wb.sheetnames:
        if net_clean in ms.strip().lower() or ms.strip().lower() in net_clean:
            return ms

    return None


def merge_network_into_master(master_wb, network_wb, network_name, log):
    """Merge a single network export into the master workbook."""
    for net_sheet_name in network_wb.sheetnames:
        net_ws = network_wb[net_sheet_name]
        net_header = find_header_row(net_ws)
        if not net_header:
            continue
        net_cols = find_network_columns(net_ws, net_header)
        if "Name" not in net_cols:
            continue

        for row in range(net_header + 1, net_ws.max_row + 1):
            name_val = net_ws.cell(row=row, column=net_cols["Name"]).value
            if not name_val:
                continue

            norm_name = normalise_name(str(name_val))
            company_val = (
                net_ws.cell(row=row, column=net_cols.get("Company", 6)).value or ""
            )

            # Determine target master sheet
            target_sheet = _resolve_target_sheet(
                master_wb, net_sheet_name, str(company_val), network_name
            )
            if not target_sheet:
                continue

            master_ws = master_wb[target_sheet]
            master_header = find_header_row(master_ws)
            if not master_header:
                continue
            master_cols = find_network_columns(master_ws, master_header)

            # Check if expert already exists
            existing_row = None
            if "Name" in master_cols:
                for mr in range(master_header + 1, master_ws.max_row + 1):
                    existing_name = master_ws.cell(
                        row=mr, column=master_cols["Name"]
                    ).value
                    if (
                        existing_name
                        and normalise_name(str(existing_name)) == norm_name
                    ):
                        existing_row = mr
                        break

            if existing_row:
                # Mark network column
                if network_name in master_cols:
                    master_ws.cell(
                        row=existing_row, column=master_cols[network_name]
                    ).value = "Yes"
                    log["networkUpdated"].append(
                        {
                            "name": str(name_val),
                            "sheet": target_sheet,
                            "network": network_name,
                        }
                    )

                # Update credits
                if "Credits" in net_cols and "Credits" in master_cols:
                    new_credit = net_ws.cell(
                        row=row, column=net_cols["Credits"]
                    ).value
                    if new_credit:
                        old_credit = master_ws.cell(
                            row=existing_row, column=master_cols["Credits"]
                        ).value
                        if old_credit and str(new_credit) != str(old_credit):
                            combined = f"{old_credit} / {new_credit}"
                            master_ws.cell(
                                row=existing_row, column=master_cols["Credits"]
                            ).value = combined
                            log["creditsUpdated"].append(
                                {
                                    "name": str(name_val),
                                    "sheet": target_sheet,
                                    "oldValue": str(old_credit),
                                    "newValue": combined,
                                }
                            )
                        elif not old_credit:
                            master_ws.cell(
                                row=existing_row, column=master_cols["Credits"]
                            ).value = new_credit
            else:
                # Insert new expert
                new_row = master_ws.max_row + 1
                copy_row_style(master_ws, new_row - 1, new_row)

                # Bullet marker - filled circle in blue
                bullet_cell = master_ws.cell(row=new_row, column=2)
                bullet_cell.value = "\u25cf"
                bullet_cell.font = openpyxl.styles.Font(color="002060", size=10)

                # Copy fields from network to master
                field_mapping = {
                    "Name": "Name",
                    "Company": "Company",
                    "Positions": "Positions",
                    "Biography": "Biography",
                    "Country": "Country",
                    "StartDate": "StartDate",
                    "EndDate": "EndDate",
                    "LinkedIn": "LinkedIn",
                }
                for net_key, master_key in field_mapping.items():
                    if net_key in net_cols and master_key in master_cols:
                        val = net_ws.cell(row=row, column=net_cols[net_key]).value
                        if val:
                            master_ws.cell(
                                row=new_row, column=master_cols[master_key]
                            ).value = val

                # Mark network column
                if network_name in master_cols:
                    master_ws.cell(
                        row=new_row, column=master_cols[network_name]
                    ).value = "Yes"

                # Credits
                if "Credits" in net_cols and "Credits" in master_cols:
                    credit_val = net_ws.cell(
                        row=row, column=net_cols["Credits"]
                    ).value
                    if credit_val:
                        master_ws.cell(
                            row=new_row, column=master_cols["Credits"]
                        ).value = credit_val

                log["added"].append(
                    {
                        "name": str(name_val),
                        "sheet": target_sheet,
                        "network": network_name,
                    }
                )


def run_merge(master_path, network_files, output_path):
    """
    Main entry point.

    master_path: path to master .xlsx
    network_files: dict of {'AlphaSights': path, 'Guidepoint': path, 'GLG': path}
    output_path: where to save the updated master

    Returns: change log dict
    """
    master_wb = openpyxl.load_workbook(master_path)

    log = {"added": [], "networkUpdated": [], "creditsUpdated": []}

    for network_name, file_path in network_files.items():
        if file_path:
            net_wb = openpyxl.load_workbook(file_path)
            merge_network_into_master(master_wb, net_wb, network_name, log)

    # Update "Last Updated" dates across sheets
    today = datetime.now()
    for sheet_name in master_wb.sheetnames:
        ws = master_wb[sheet_name]
        for col in range(20, 30):
            cell_val = ws.cell(row=2, column=col).value
            if cell_val and "last updated" in str(cell_val).lower():
                ws.cell(row=2, column=col + 1).value = today
                break

    master_wb.save(output_path)
    return log


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Merge network exports into master tracker")
    parser.add_argument("master", help="Path to master tracker .xlsx")
    parser.add_argument("output", help="Path to save updated master .xlsx")
    parser.add_argument("--alphasights", help="Path to AlphaSights export .xlsx")
    parser.add_argument("--guidepoint", help="Path to Guidepoint export .xlsx")
    parser.add_argument("--glg", help="Path to GLG export .xlsx")
    args = parser.parse_args()

    network_files = {}
    if args.alphasights:
        network_files["AlphaSights"] = args.alphasights
    if args.guidepoint:
        network_files["Guidepoint"] = args.guidepoint
    if args.glg:
        network_files["GLG"] = args.glg

    if not network_files:
        print(json.dumps({"error": "No network files provided"}))
        sys.exit(1)

    log = run_merge(args.master, network_files, args.output)
    print(json.dumps(log))
